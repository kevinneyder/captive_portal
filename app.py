from flask import Flask, request, render_template, redirect
from openpyxl import Workbook, load_workbook
from datetime import datetime
import requests
import os

app = Flask(__name__)

# ==============================
# CONFIGURATION
# ==============================
EXCEL_FILE = "happy_tea_clients.xlsx"
UNIFI_CONTROLLER_IP = "192.168.100.196"
UNIFI_PORT = 8443
UNIFI_USERNAME = "admin"
UNIFI_PASSWORD = "admin"
UNIFI_SITE = "default"
DEFAULT_REDIRECT = "http://connectivitycheck.gstatic.com/generate_204"

requests.packages.urllib3.disable_warnings()


# ==============================
# UTILITIES
# ==============================
def log(msg: str, level: str = "INFO"):
    """Print formatted log messages to console."""
    print(f"[{level}] {msg}")


def safe_get(data: dict, key: str, default=""):
    """Safely get a value from a dictionary and trim spaces."""
    return data.get(key, "").strip() or default


# ==============================
# EXCEL HANDLING
# ==============================
def init_excel():
    """Create Excel file if it doesn’t exist."""
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.append([
            "Age", "Gender", "Zone", "Source Info", "Favorite Drink",
            "Register Date", "Client MAC", "User Agent"
        ])
        wb.save(EXCEL_FILE)
        log("Excel file created.")


def save_to_excel(age, gender, zone, sourceInfo, drink, client_id, user_agent):
    """Save form data to the Excel file."""
    wb = load_workbook(EXCEL_FILE) if os.path.exists(EXCEL_FILE) else Workbook()
    ws = wb.active

    # Ensure headers
    if ws.max_row == 1:
        ws.append([
            "Age", "Gender", "Zone", "Source Info", "Favorite Drink",
            "Register Date", "Client MAC", "User Agent"
        ])

    ws.append([
        age, gender, zone, sourceInfo, drink,
        datetime.now().strftime("%Y-%m-%d %H:%M"),
        client_id, user_agent
    ])
    wb.save(EXCEL_FILE)
    log("Data saved to Excel.", "OK")


# ==============================
# UNIFI API
# ==============================
def unifi_login():
    """Log in to UniFi Controller and return the session."""
    session = requests.Session()
    session.verify = False
    login_url = f"https://{UNIFI_CONTROLLER_IP}:{UNIFI_PORT}/api/login"

    try:
        response = session.post(login_url, json={
            "username": UNIFI_USERNAME,
            "password": UNIFI_PASSWORD
        }, timeout=10)

        if response.ok:
            log("UniFi session started successfully.", "OK")
            return session

        log(f"Login failed: {response.status_code} - {response.text}", "ERROR")
    except Exception as e:
        log(f"Could not connect to UniFi Controller: {e}", "ERROR")

    return None


def unifi_authorize_guest(session, client_mac, minutes=1440):
    """Authorize a guest client through UniFi API."""
    url = f"https://{UNIFI_CONTROLLER_IP}:{UNIFI_PORT}/api/s/{UNIFI_SITE}/cmd/stamgr"
    data = {"cmd": "authorize-guest", "mac": client_mac, "minutes": minutes}

    try:
        response = session.post(url, json=data, timeout=10)
        if response.ok:
            log(f"Client {client_mac} authorized for {minutes} minutes.", "OK")
            return True

        log(f"Authorization failed: {response.status_code} - {response.text}", "ERROR")
    except Exception as e:
        log(f"Error authorizing client: {e}", "ERROR")

    return False


# ==============================
# ROUTES
# ==============================
@app.route("/")
@app.route("/guest/")
@app.route("/guest/s/default/")
def portal():
    """Main captive portal page."""
    params = {
        "client_id": request.args.get("id"),
        "ap_mac": request.args.get("ap"),
        "ssid": request.args.get("ssid"),
        "redirect_url": request.args.get("url", DEFAULT_REDIRECT),
        "user_agent": request.headers.get("User-Agent", "")
    }

    log("New UniFi request:")
    for k, v in params.items():
        log(f"  {k}: {v}")

    return render_template("portal.html", **params)


@app.route("/save", methods=["POST"])
def save():
    """Handle form submission and authorize guest in UniFi."""
    try:
        form = request.form
        age = safe_get(form, "age")
        gender = safe_get(form, "gender")
        zone = safe_get(form, "zone")
        sourceInfo = safe_get(form, "sourceInfo")
        drink = safe_get(form, "drink")
        client_id = safe_get(form, "client_id")
        redirect_url = safe_get(form, "redirect_url", DEFAULT_REDIRECT)
        user_agent = request.headers.get("User-Agent", "")

        if not all([age, gender, zone, sourceInfo, drink]):
            return "Missing required fields", 400

        # Save data to Excel
        save_to_excel(age, gender, zone, sourceInfo, drink, client_id, user_agent)

        # Authorize client in UniFi
        if client_id:
            log(f"Attempting to authorize client: {client_id}")
            session = unifi_login()
            if session and unifi_authorize_guest(session, client_id):
                log("Client authorized successfully via UniFi API.", "OK")
            else:
                log("Client could not be authorized.", "ERROR")
        else:
            log("No 'client_id' received in form.", "WARN")

        log("Redirecting to Google.com")
        return redirect("https://www.google.com")

    except Exception as e:
        log(f"Internal error: {e}", "ERROR")
        return f"Internal error: {e}", 500


# ==============================
# CAPTIVE PORTAL DETECTION
# ==============================
@app.route("/generate_204")
@app.route("/hotspot-detect.html")
@app.route("/library/test/success.html")
def captive_endpoints():
    """Endpoints used by Android/iOS to detect captive portals."""
    return "", 204


@app.route("/ncsi.txt")
def ncsi():
    """Endpoint used by Windows for connectivity check."""
    return "Microsoft NCSI"


# ==============================
# MAIN ENTRY POINT
# ==============================
if __name__ == "__main__":
    init_excel()
    app.run(host="0.0.0.0", port=80, debug=True)
